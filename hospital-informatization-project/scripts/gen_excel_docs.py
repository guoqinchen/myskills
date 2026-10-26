#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Excel文档批量生成脚本
生成：需求调研表、费用明细表、评分标准、验收检查表
"""

import argparse
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 样式定义
TITLE_FONT = Font(name="仿宋", size=14, bold=True)
HEADER_FONT = Font(name="仿宋", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SUB_HEADER_FONT = Font(name="仿宋", size=10, bold=True)
SUB_HEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
NORMAL_FONT = Font(name="仿宋", size=10)
BOLD_FONT = Font(name="仿宋", size=10, bold=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
NOTE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")


def apply_header_style(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 25


def apply_data_style(ws, row, col_count, fill=None):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = NORMAL_FONT
        cell.alignment = LEFT_ALIGN
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill


def gen_requirements_survey(project_name, output_dir):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "功能需求调研"
    ws1.merge_cells("A1:K1")
    ws1["A1"] = f"{project_name}需求调研表 - 功能需求"
    ws1["A1"].font = TITLE_FONT
    ws1["A1"].alignment = CENTER_ALIGN
    ws1.row_dimensions[1].height = 30

    headers = ["序号", "需求项编号", "模块名称", "需求名称", "需求描述", "需求类型", "优先级", "使用角色", "数据来源", "来源", "备注"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=2, column=col, value=h)
    apply_header_style(ws1, 2, len(headers))

    for row in range(3, 13):
        ws1.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws1, row, len(headers))

    widths = [6, 14, 15, 18, 35, 10, 8, 15, 14, 12, 20]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A3"

    ws2 = wb.create_sheet("数据需求调研")
    ws2.merge_cells("A1:J1")
    ws2["A1"] = f"{project_name}需求调研表 - 数据需求"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = CENTER_ALIGN
    headers2 = ["序号", "数据项编号", "数据项名称", "数据类型", "数据长度", "数据来源", "是否必填", "校验规则", "关联实体", "备注"]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=2, column=col, value=h)
    apply_header_style(ws2, 2, len(headers2))
    for row in range(3, 13):
        ws2.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws2, row, len(headers2))
    widths2 = [6, 14, 16, 10, 10, 14, 10, 20, 15, 20]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A3"

    ws3 = wb.create_sheet("接口需求调研")
    ws3.merge_cells("A1:J1")
    ws3["A1"] = f"{project_name}需求调研表 - 接口需求"
    ws3["A1"].font = TITLE_FONT
    ws3["A1"].alignment = CENTER_ALIGN
    headers3 = ["序号", "接口编号", "接口名称", "对接系统", "接口方向", "数据内容", "接口协议", "数据格式", "调用频率", "备注"]
    for col, h in enumerate(headers3, 1):
        ws3.cell(row=2, column=col, value=h)
    apply_header_style(ws3, 2, len(headers3))
    for row in range(3, 10):
        ws3.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws3, row, len(headers3))
    widths3 = [6, 14, 18, 16, 10, 20, 12, 12, 12, 20]
    for i, w in enumerate(widths3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A3"

    ws4 = wb.create_sheet("安全需求调研")
    ws4.merge_cells("A1:H1")
    ws4["A1"] = f"{project_name}需求调研表 - 安全需求"
    ws4["A1"].font = TITLE_FONT
    ws4["A1"].alignment = CENTER_ALIGN
    headers4 = ["序号", "安全项编号", "安全需求", "安全类别", "合规依据", "实现方式", "优先级", "备注"]
    for col, h in enumerate(headers4, 1):
        ws4.cell(row=2, column=col, value=h)
    apply_header_style(ws4, 2, len(headers4))
    for row in range(3, 10):
        ws4.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws4, row, len(headers4))
    widths4 = [6, 14, 25, 15, 20, 25, 8, 20]
    for i, w in enumerate(widths4, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = "A3"

    ws5 = wb.create_sheet("用户角色调研")
    ws5.merge_cells("A1:I1")
    ws5["A1"] = f"{project_name}需求调研表 - 用户角色"
    ws5["A1"].font = TITLE_FONT
    ws5["A1"].alignment = CENTER_ALIGN
    headers5 = ["序号", "角色编号", "角色名称", "角色描述", "用户数量", "权限范围", "数据权限", "关联部门", "备注"]
    for col, h in enumerate(headers5, 1):
        ws5.cell(row=2, column=col, value=h)
    apply_header_style(ws5, 2, len(headers5))
    for row in range(3, 10):
        ws5.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws5, row, len(headers5))
    widths5 = [6, 12, 14, 25, 10, 25, 20, 14, 20]
    for i, w in enumerate(widths5, 1):
        ws5.column_dimensions[get_column_letter(i)].width = w
    ws5.freeze_panes = "A3"

    filepath = os.path.join(output_dir, f"{project_name}需求调研表.xlsx")
    wb.save(filepath)
    print(f"需求调研表已生成: {filepath}")


def gen_cost_detail(project_name, output_dir, cost_items=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "费用明细"
    ws.merge_cells("A1:D1")
    ws["A1"] = f"{project_name}费用明细表"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 30

    headers = ["序号", "功能名称", "功能参数", "报价（万元）"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=h)
    apply_header_style(ws, 2, len(headers))

    if cost_items is None:
        cost_items = [
            ("基础信息管理", "机构信息维护、科室架构管理、人员信息管理、数据字典维护", 0.5),
            ("用户与权限管理", "角色定义、权限分配、操作日志、登录管理", 1.0),
            ("工作台", "待办提醒、消息通知、数据概览、快捷入口", 1.0),
            ("数据采集与录入", "表单设计、批量导入、数据校验、附件上传", 2.0),
            ("业务流程管理", "流程配置、流程审批、流程监控、流程统计", 3.0),
            ("统计报表", "自定义报表、图表展示、数据导出、定期推送", 2.0),
            ("数据分析中心", "多维度分析、趋势分析、对比分析、预警分析", 2.0),
            ("系统接口对接", "与HIS/EMR/LIS等系统数据交换接口", 1.0),
            ("移动端应用", "移动端查询、移动端审批、消息推送", 2.0),
            ("安全管理", "数据加密、访问控制、审计日志、备份恢复", 1.0),
            ("服务器硬件", "6核以上CPU，32G内存，2T硬盘，本地部署", 1.0),
            ("安装及部署", "系统安装、环境配置、数据初始化、联调测试", 0.5),
        ]

    for idx, (name, desc, price) in enumerate(cost_items, 1):
        row = idx + 2
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=3, value=desc)
        ws.cell(row=row, column=4, value=price)
        apply_data_style(ws, row, len(headers))
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).alignment = CENTER_ALIGN if col != 3 else LEFT_ALIGN

    total_row = len(cost_items) + 3
    ws.cell(row=total_row, column=1, value="合计")
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    ws.cell(row=total_row, column=4, value=f"=SUM(D3:D{total_row - 1})")
    apply_data_style(ws, total_row, len(headers), TOTAL_FILL)
    ws.cell(row=total_row, column=1).font = BOLD_FONT
    ws.cell(row=total_row, column=1).alignment = CENTER_ALIGN
    ws.cell(row=total_row, column=4).font = BOLD_FONT
    ws.cell(row=total_row, column=4).alignment = CENTER_ALIGN

    note_row = total_row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    note_text = (
        f"备注：系统不含第三方接口费用，如HIS等系统接口费用，如需接入另行计算。"
        f"一年内提供免费远程技术服务。系统应用场景说明：本系统适用于{project_name}相关场景。"
    )
    ws.cell(row=note_row, column=1, value=note_text)
    ws.cell(row=note_row, column=1).font = NORMAL_FONT
    ws.cell(row=note_row, column=1).alignment = LEFT_ALIGN
    ws.cell(row=note_row, column=1).fill = NOTE_FILL
    ws.cell(row=note_row, column=1).border = THIN_BORDER
    ws.row_dimensions[note_row].height = 45

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 12

    filepath = os.path.join(output_dir, f"{project_name}费用明细表.xlsx")
    wb.save(filepath)
    print(f"费用明细表已生成: {filepath}")


def gen_scoring_standard(project_name, output_dir):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "评分总表"
    ws1.merge_cells("A1:D1")
    ws1["A1"] = f"{project_name}招标采购评分标准"
    ws1["A1"].font = TITLE_FONT
    ws1["A1"].alignment = CENTER_ALIGN
    ws1.row_dimensions[1].height = 30
    headers = ["评分类别", "分值", "权重", "说明"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=2, column=col, value=h)
    apply_header_style(ws1, 2, len(headers))
    categories = [
        ("价格评分", 30, "30%", "满足招标要求且最低评标价为基准价，低价优先法"),
        ("功能评分", 50, "50%", "技术方案、功能响应、演示效果、安全方案、实施方案"),
        ("商务评分", 20, "20%", "企业资质、类似业绩、售后服务、团队配置、本地化服务"),
        ("合计", 100, "100%", ""),
    ]
    for idx, (cat, score, weight, desc) in enumerate(categories, 3):
        ws1.cell(row=idx, column=1, value=cat)
        ws1.cell(row=idx, column=2, value=score)
        ws1.cell(row=idx, column=3, value=weight)
        ws1.cell(row=idx, column=4, value=desc)
        fill = TOTAL_FILL if cat == "合计" else None
        apply_data_style(ws1, idx, len(headers), fill)
        if cat == "合计":
            for col in range(1, len(headers) + 1):
                ws1.cell(row=idx, column=col).font = BOLD_FONT
    ws1.column_dimensions["A"].width = 15
    ws1.column_dimensions["B"].width = 8
    ws1.column_dimensions["C"].width = 8
    ws1.column_dimensions["D"].width = 50

    ws2 = wb.create_sheet("价格评分(30分)")
    ws2.merge_cells("A1:C1")
    ws2["A1"] = "价格评分标准（满分30分）"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = CENTER_ALIGN
    for col, h in enumerate(["评分项", "分值", "评分标准"], 1):
        ws2.cell(row=2, column=col, value=h)
    apply_header_style(ws2, 2, 3)
    ws2.cell(row=3, column=1, value="投标报价")
    ws2.cell(row=3, column=2, value=30)
    ws2.cell(row=3, column=3, value="价格分=（基准价/投标报价）x30，四舍五入保留两位小数。基准价为满足招标文件要求且投标价格最低的报价。超过预算限价的报价为无效投标。")
    apply_data_style(ws2, 3, 3)
    ws2.column_dimensions["A"].width = 15
    ws2.column_dimensions["B"].width = 8
    ws2.column_dimensions["C"].width = 60

    ws3 = wb.create_sheet("功能评分(50分)")
    ws3.merge_cells("A1:C1")
    ws3["A1"] = "功能评分标准（满分50分）"
    ws3["A1"].font = TITLE_FONT
    ws3["A1"].alignment = CENTER_ALIGN
    for col, h in enumerate(["评分项", "分值", "评分标准"], 1):
        ws3.cell(row=2, column=col, value=h)
    apply_header_style(ws3, 2, 3)
    func_items = [
        ("总体技术方案", 8, "方案完整性、合理性、先进性。优8-7分，良6-5分，中4-3分，差2-0分"),
        ("系统架构设计", 6, "架构合理性、可扩展性、安全性。优6-5分，良4-3分，差2-0分"),
        ("功能响应度", 15, "完全响应得满分，每偏离一项扣1分。带星号关键需求不响应视为无效投标"),
        ("系统演示/原型", 8, "有完整演示得8分，有部分演示得4分，无演示得0分"),
        ("数据安全方案", 5, "方案完整、符合等保要求得5分，基本满足得3分，不满足得0分"),
        ("实施方案", 4, "实施计划合理、风险控制完善。优4分，良3分，中2分，差0-1分"),
        ("培训方案", 4, "培训计划完整、内容针对性强。优4分，良3分，中2分，差0-1分"),
    ]
    for idx, (item, score, std) in enumerate(func_items, 3):
        ws3.cell(row=idx, column=1, value=item)
        ws3.cell(row=idx, column=2, value=score)
        ws3.cell(row=idx, column=3, value=std)
        apply_data_style(ws3, idx, 3)
    total_row = len(func_items) + 3
    ws3.cell(row=total_row, column=1, value="合计")
    ws3.cell(row=total_row, column=2, value=50)
    ws3.cell(row=total_row, column=3, value="")
    apply_data_style(ws3, total_row, 3, TOTAL_FILL)
    for col in range(1, 4):
        ws3.cell(row=total_row, column=col).font = BOLD_FONT
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 8
    ws3.column_dimensions["C"].width = 60

    ws4 = wb.create_sheet("商务评分(20分)")
    ws4.merge_cells("A1:C1")
    ws4["A1"] = "商务评分标准（满分20分）"
    ws4["A1"].font = TITLE_FONT
    ws4["A1"].alignment = CENTER_ALIGN
    for col, h in enumerate(["评分项", "分值", "评分标准"], 1):
        ws4.cell(row=2, column=col, value=h)
    apply_header_style(ws4, 2, 3)
    biz_items = [
        ("企业资质", 4, "ISO9001认证得1分；ISO27001认证得1分；CMMI3级及以上得2分"),
        ("类似业绩", 6, "每提供一个同类项目合同得2分，满分6分（需提供合同复印件及验收证明）"),
        ("售后服务", 5, "免费维护期>=1年得2分，每增加1年加1分（满分4分）；7x24小时响应得1分"),
        ("团队配置", 3, "项目经理PMP认证得1分；团队具有相关技术认证每人得0.5分（满分2分）"),
        ("本地化服务", 2, "在本地有常驻服务机构得2分，承诺提供本地化服务得1分"),
    ]
    for idx, (item, score, std) in enumerate(biz_items, 3):
        ws4.cell(row=idx, column=1, value=item)
        ws4.cell(row=idx, column=2, value=score)
        ws4.cell(row=idx, column=3, value=std)
        apply_data_style(ws4, idx, 3)
    total_row4 = len(biz_items) + 3
    ws4.cell(row=total_row4, column=1, value="合计")
    ws4.cell(row=total_row4, column=2, value=20)
    ws4.cell(row=total_row4, column=3, value="")
    apply_data_style(ws4, total_row4, 3, TOTAL_FILL)
    for col in range(1, 4):
        ws4.cell(row=total_row4, column=col).font = BOLD_FONT
    ws4.column_dimensions["A"].width = 15
    ws4.column_dimensions["B"].width = 8
    ws4.column_dimensions["C"].width = 60

    ws5 = wb.create_sheet("评分汇总表")
    ws5.merge_cells("A1:F1")
    ws5["A1"] = "评分汇总表"
    ws5["A1"].font = TITLE_FONT
    ws5["A1"].alignment = CENTER_ALIGN
    for col, h in enumerate(["投标人", "价格得分", "功能得分", "商务得分", "总分", "排名"], 1):
        ws5.cell(row=2, column=col, value=h)
    apply_header_style(ws5, 2, 6)
    for row in range(3, 6):
        ws5.cell(row=row, column=1, value=f"供应商{chr(64 + row - 2)}")
        apply_data_style(ws5, row, 6)
    ws5.column_dimensions["A"].width = 15
    for col in range(2, 7):
        ws5.column_dimensions[get_column_letter(col)].width = 12

    filepath = os.path.join(output_dir, f"{project_name}评分标准.xlsx")
    wb.save(filepath)
    print(f"评分标准表已生成: {filepath}")


def gen_acceptance_checklist(project_name, output_dir):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "功能验收检查"
    ws1.merge_cells("A1:H1")
    ws1["A1"] = f"{project_name}功能验收检查表"
    ws1["A1"].font = TITLE_FONT
    ws1["A1"].alignment = CENTER_ALIGN
    ws1.row_dimensions[1].height = 30
    headers = ["序号", "验收项编号", "功能模块", "验收项名称", "验收标准", "验收方法", "验收结果", "问题描述"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=2, column=col, value=h)
    apply_header_style(ws1, 2, len(headers))
    for row in range(3, 13):
        ws1.cell(row=row, column=1, value=row - 2)
        apply_data_style(ws1, row, len(headers))
    widths = [6, 14, 15, 18, 30, 14, 12, 25]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A3"

    ws2 = wb.create_sheet("技术指标验收")
    ws2.merge_cells("A1:G1")
    ws2["A1"] = f"{project_name}技术指标验收表"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = CENTER_ALIGN
    headers2 = ["序号", "指标名称", "合同要求值", "实测值", "是否达标", "测试方法", "备注"]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=2, column=col, value=h)
    apply_header_style(ws2, 2, len(headers2))
    tech_items = [
        ("系统响应时间", "<=3秒", "", "", "页面操作计时", ""),
        ("并发用户数", ">=100", "", "", "压力测试", ""),
        ("系统可用率", ">=99.5%", "", "", "运行日志统计", ""),
        ("数据加密存储", "AES-256", "", "", "安全扫描", ""),
        ("等保级别", "二级", "", "", "等保测评报告", ""),
        ("数据备份", "每日自动备份", "", "", "备份日志检查", ""),
    ]
    for idx, (name, req, measured, pass_flag, method, note) in enumerate(tech_items, 3):
        ws2.cell(row=idx, column=1, value=idx - 2)
        ws2.cell(row=idx, column=2, value=name)
        ws2.cell(row=idx, column=3, value=req)
        ws2.cell(row=idx, column=4, value=measured)
        ws2.cell(row=idx, column=5, value=pass_flag)
        ws2.cell(row=idx, column=6, value=method)
        ws2.cell(row=idx, column=7, value=note)
        apply_data_style(ws2, idx, len(headers2))
    widths2 = [6, 16, 16, 14, 10, 18, 20]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A3"

    ws3 = wb.create_sheet("交付物验收")
    ws3.merge_cells("A1:E1")
    ws3["A1"] = f"{project_name}交付物验收表"
    ws3["A1"].font = TITLE_FONT
    ws3["A1"].alignment = CENTER_ALIGN
    headers3 = ["序号", "交付物名称", "合同要求", "交付情况", "质量评定"]
    for col, h in enumerate(headers3, 1):
        ws3.cell(row=2, column=col, value=h)
    apply_header_style(ws3, 2, len(headers3))
    deliverables = [
        ("系统软件（生产环境部署）", "完整部署并通过验收测试", "", ""),
        ("需求详细说明书", "经甲方确认签字", "", ""),
        ("系统开发文档", "含架构设计、数据库设计、接口设计", "", ""),
        ("系统操作手册", "覆盖全部功能模块", "", ""),
        ("培训记录", "含培训签到表、培训照片", "", ""),
        ("系统源代码", "含注释、编译说明", "", ""),
        ("测试报告", "含功能测试、性能测试、安全测试", "", ""),
        ("验收报告", "双方签字盖章", "", ""),
    ]
    for idx, (name, req, delivered, quality) in enumerate(deliverables, 3):
        ws3.cell(row=idx, column=1, value=idx - 2)
        ws3.cell(row=idx, column=2, value=name)
        ws3.cell(row=idx, column=3, value=req)
        ws3.cell(row=idx, column=4, value=delivered)
        ws3.cell(row=idx, column=5, value=quality)
        apply_data_style(ws3, idx, len(headers3))
    widths3 = [6, 25, 30, 12, 12]
    for i, w in enumerate(widths3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A3"

    filepath = os.path.join(output_dir, f"{project_name}验收检查表.xlsx")
    wb.save(filepath)
    print(f"验收检查表已生成: {filepath}")


def main():
    parser = argparse.ArgumentParser(description="生成项目Excel文档")
    parser.add_argument("--project", required=True, help="项目名称")
    parser.add_argument("--type", default="all", choices=["all", "survey", "cost", "scoring", "acceptance"],
                        help="生成类型")
    parser.add_argument("--output", required=True, help="输出目录")
    args = parser.parse_args()
    os.makedirs(args.output, exist_ok=True)

    if args.type in ("all", "survey"):
        gen_requirements_survey(args.project, args.output)
    if args.type in ("all", "cost"):
        gen_cost_detail(args.project, args.output)
    if args.type in ("all", "scoring"):
        gen_scoring_standard(args.project, args.output)
    if args.type in ("all", "acceptance"):
        gen_acceptance_checklist(args.project, args.output)
    print(f"\n所有Excel文档已生成到: {args.output}")


if __name__ == "__main__":
    main()
