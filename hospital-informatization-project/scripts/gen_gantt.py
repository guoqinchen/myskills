#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
项目实施甘特图生成脚本
生成包含标准任务列表和甘特条形图的Excel文件
"""

import argparse
import os
import sys
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

# 标准任务模板
STANDARD_TASKS = [
    # (任务名称, 阶段, 工期天数, 责任人, 阶段颜色)
    ("需求调研阶段", "", 0, "", ""),
    ("项目启动会", "需求调研", 1, "项目组全体", "4472C4"),
    ("需求调研访谈", "需求调研", 3, "需求分析师", "4472C4"),
    ("需求文档编写", "需求调研", 2, "需求分析师", "4472C4"),
    ("需求评审确认", "需求调研", 1, "项目组+用户", "4472C4"),
    ("设计阶段", "", 0, "", ""),
    ("系统架构设计", "设计", 2, "架构师", "ED7D31"),
    ("数据库设计", "设计", 2, "DBA", "ED7D31"),
    ("界面原型设计", "设计", 3, "UI设计师", "ED7D31"),
    ("设计评审", "设计", 1, "项目组+用户", "ED7D31"),
    ("开发阶段", "", 0, "", ""),
    ("环境搭建", "开发", 1, "运维工程师", "A5A5A5"),
    ("基础模块开发", "开发", 5, "开发团队", "A5A5A5"),
    ("业务模块开发", "开发", 8, "开发团队", "A5A5A5"),
    ("接口开发", "开发", 3, "开发团队", "A5A5A5"),
    ("单元测试", "开发", 2, "开发团队", "A5A5A5"),
    ("测试阶段", "", 0, "", ""),
    ("集成测试", "测试", 2, "测试工程师", "FFC000"),
    ("系统测试", "测试", 3, "测试工程师", "FFC000"),
    ("用户验收测试", "测试", 2, "用户+测试", "FFC000"),
    ("性能测试", "测试", 1, "测试工程师", "FFC000"),
    ("安全测试", "测试", 1, "安全工程师", "FFC000"),
    ("部署上线阶段", "", 0, "", ""),
    ("生产环境部署", "部署上线", 1, "运维工程师", "5B9BD5"),
    ("数据迁移", "部署上线", 1, "DBA+开发", "5B9BD5"),
    ("用户培训", "部署上线", 2, "培训师", "5B9BD5"),
    ("试运行", "部署上线", 5, "项目组+用户", "5B9BD5"),
    ("验收阶段", "", 0, "", ""),
    ("验收准备", "验收", 2, "项目组", "70AD47"),
    ("正式验收", "验收", 1, "验收组", "70AD47"),
    ("项目交接", "验收", 1, "项目组", "70AD47"),
]

# 样式定义
HEADER_FONT = Font(name="仿宋", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
NORMAL_FONT = Font(name="仿宋", size=10)
BOLD_FONT = Font(name="仿宋", size=10, bold=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def generate_gantt(project_name, start_date_str, weeks, output_path):
    """生成甘特图Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "甘特图"

    start_date = datetime.strptime(start_date_str, "%Y-%m-%d")
    total_days = weeks * 7

    # 标题行
    ws.merge_cells("A1:I1")
    ws["A1"] = f"{project_name}项目实施甘特图"
    ws["A1"].font = Font(name="仿宋", size=14, bold=True)
    ws["A1"].alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 30

    # 表头
    headers = ["序号", "任务名称", "阶段", "开始日期", "结束日期", "工期(天)", "责任人", "状态", "进度(%)"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[2].height = 25

    # 甘特图日期列
    gantt_start_col = 10  # J列开始为甘特图区域
    for day in range(total_days):
        col = gantt_start_col + day
        current_date = start_date + timedelta(days=day)
        cell = ws.cell(row=2, column=col, value=current_date.strftime("%m/%d"))
        cell.font = Font(name="仿宋", size=8, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", textRotation=90)
        cell.fill = HEADER_FILL
        cell.font = Font(name="仿宋", size=8, bold=True, color="FFFFFF")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = 3.5

    # 填充任务数据
    row = 3
    task_seq = 0
    current_date_offset = 0

    for task_name, phase, duration, owner, color in STANDARD_TASKS:
        if duration == 0:
            # 阶段标题行
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9 + total_days)
            cell = ws.cell(row=row, column=1, value=task_name)
            cell.font = Font(name="仿宋", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
            cell.alignment = LEFT_ALIGN
            cell.border = THIN_BORDER
            ws.row_dimensions[row].height = 22
            row += 1
            continue

        task_seq += 1
        task_start = start_date + timedelta(days=current_date_offset)
        task_end = task_start + timedelta(days=duration - 1)

        # 填充基本信息
        ws.cell(row=row, column=1, value=task_seq).font = NORMAL_FONT
        ws.cell(row=row, column=1).alignment = CENTER_ALIGN

        ws.cell(row=row, column=2, value=task_name).font = NORMAL_FONT
        ws.cell(row=row, column=2).alignment = LEFT_ALIGN

        ws.cell(row=row, column=3, value=phase).font = NORMAL_FONT
        ws.cell(row=row, column=3).alignment = CENTER_ALIGN

        ws.cell(row=row, column=4, value=task_start.strftime("%Y-%m-%d")).font = NORMAL_FONT
        ws.cell(row=row, column=4).alignment = CENTER_ALIGN

        ws.cell(row=row, column=5, value=task_end.strftime("%Y-%m-%d")).font = NORMAL_FONT
        ws.cell(row=row, column=5).alignment = CENTER_ALIGN

        ws.cell(row=row, column=6, value=duration).font = NORMAL_FONT
        ws.cell(row=row, column=6).alignment = CENTER_ALIGN

        ws.cell(row=row, column=7, value=owner).font = NORMAL_FONT
        ws.cell(row=row, column=7).alignment = CENTER_ALIGN

        ws.cell(row=row, column=8, value="未开始").font = NORMAL_FONT
        ws.cell(row=row, column=8).alignment = CENTER_ALIGN

        ws.cell(row=row, column=9, value=0).font = NORMAL_FONT
        ws.cell(row=row, column=9).alignment = CENTER_ALIGN
        ws.cell(row=row, column=9).number_format = "0%"

        # 绘制甘特条
        for day_offset in range(duration):
            gantt_col = gantt_start_col + current_date_offset + day_offset
            cell = ws.cell(row=row, column=gantt_col, value="")
            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            cell.border = THIN_BORDER

        # 设置边框
        for col in range(1, 10):
            ws.cell(row=row, column=col).border = THIN_BORDER

        ws.row_dimensions[row].height = 20
        current_date_offset += duration
        row += 1

    # 设置列宽
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 13
    ws.column_dimensions["F"].width = 9
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 9

    # 冻结窗格
    ws.freeze_panes = "C3"

    # 添加说明Sheet
    ws2 = wb.create_sheet("说明")
    ws2["A1"] = "甘特图使用说明"
    ws2["A1"].font = Font(name="仿宋", size=14, bold=True)
    explanations = [
        "",
        "1. 本甘特图按标准项目实施流程生成，包含6个阶段：需求调研、设计、开发、测试、部署上线、验收",
        "2. 各阶段任务可根据实际项目需求调整工期和责任人",
        "3. 甘特条颜色区分：蓝色=需求调研，橙色=设计，灰色=开发，黄色=测试，浅蓝=部署，绿色=验收",
        "4. 状态列可手动更新为：未开始/进行中/已完成/已延期",
        "5. 进度列可手动更新百分比（0%-100%）",
        "6. 开始日期和结束日期可调整，甘特条需同步更新",
        "",
        f"项目总工期：{current_date_offset}天（约{current_date_offset/7:.1f}周）",
        f"项目开始日期：{start_date_str}",
        f"预计结束日期：{(start_date + timedelta(days=current_date_offset-1)).strftime('%Y-%m-%d')}",
    ]
    for i, text in enumerate(explanations, 2):
        ws2.cell(row=i, column=1, value=text).font = Font(name="仿宋", size=11)
    ws2.column_dimensions["A"].width = 80

    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    wb.save(output_path)
    print(f"甘特图已生成: {output_path}")
    print(f"项目总工期: {current_date_offset}天（约{current_date_offset/7:.1f}周）")
    print(f"任务总数: {task_seq}")


def main():
    parser = argparse.ArgumentParser(description="生成项目实施甘特图Excel文件")
    parser.add_argument("--project", required=True, help="项目名称")
    parser.add_argument("--start", default="2026-01-01", help="开始日期 (YYYY-MM-DD)")
    parser.add_argument("--weeks", type=int, default=12, help="甘特图总周数")
    parser.add_argument("--output", required=True, help="输出文件路径")

    args = parser.parse_args()
    generate_gantt(args.project, args.start, args.weeks, args.output)


if __name__ == "__main__":
    main()
